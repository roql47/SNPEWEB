import openpyxl
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

wb = openpyxl.load_workbook(r'c:\Users\User\Desktop\SNPEweb2\브랜드 홈페이지 개설안.xlsx')
print("=== Sheets ===")
print(wb.sheetnames)

for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n{'='*60}")
    print(f"Sheet: {name}")
    print(f"Rows: {ws.max_row}, Cols: {ws.max_column}")
    print(f"{'='*60}")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column, values_only=False):
        values = []
        for cell in row:
            v = cell.value
            if v is not None:
                s = str(v).replace('\n', ' ').strip()
                if s:
                    values.append(f"[{cell.coordinate}] {s}")
        if values:
            print(" | ".join(values))
